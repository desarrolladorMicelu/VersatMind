"""
Generador de informes PDF y Excel para Mind by Versat.
Incluye branding del cliente configurable.
Requisitos: 6.1, 6.2, 6.3, 6.7, 6.8
"""
from __future__ import annotations

import asyncio
import os
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Literal


class ReportGenerationError(Exception):
    """Se lanza cuando falla la generación del informe."""
    pass


async def generate_report(
    data: list[dict[str, Any]],
    format: Literal["pdf", "excel"],
    title: str,
    period: str,
    client_name: str = "Versat",
    client_logo_path: str | None = None,
) -> str:
    """
    Genera un informe PDF o Excel y retorna la ruta absoluta del archivo temporal.

    El llamador (Orchestrator/Handler) es responsable de eliminar el archivo
    después del envío.

    Timeout: 60 segundos. Soporta hasta 10.000 filas.
    Requisitos: 6.1, 6.2, 6.3, 6.7, 6.8
    """
    if not data:
        raise ReportGenerationError("No hay datos para generar el informe.")

    try:
        async with asyncio.timeout(60):
            if format == "pdf":
                return await asyncio.to_thread(
                    _generate_pdf, data, title, period, client_name, client_logo_path
                )
            elif format == "excel":
                return await asyncio.to_thread(
                    _generate_excel, data, title, period, client_name
                )
            else:
                raise ReportGenerationError(f"Formato no soportado: {format!r}")
    except asyncio.TimeoutError as exc:
        raise ReportGenerationError("Timeout: la generación del informe excedió 60 segundos.") from exc
    except ReportGenerationError:
        raise
    except Exception as exc:
        raise ReportGenerationError(f"Error al generar informe: {exc}") from exc


def _generate_pdf(
    data: list[dict[str, Any]],
    title: str,
    period: str,
    client_name: str,
    client_logo_path: str | None,
) -> str:
    """Genera PDF con ReportLab. Retorna ruta del archivo temporal."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    )

    tmp_dir = tempfile.mkdtemp(prefix="mind_report_")
    safe_title = "".join(c if c.isalnum() or c in " _-" else "_" for c in title)
    file_path = os.path.join(tmp_dir, f"{safe_title}_{period.replace(' ', '_')}.pdf")

    page_size = landscape(A4) if len(data) > 0 and len(data[0]) > 6 else A4
    doc = SimpleDocTemplate(
        file_path,
        pagesize=page_size,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    styles = getSampleStyleSheet()
    header_style = ParagraphStyle(
        "header", parent=styles["Heading1"],
        fontSize=16, spaceAfter=4, textColor=colors.HexColor("#1a1a2e")
    )
    meta_style = ParagraphStyle(
        "meta", parent=styles["Normal"],
        fontSize=9, textColor=colors.grey
    )

    elements = []

    # --- Encabezado ---
    if client_logo_path and Path(client_logo_path).exists():
        try:
            logo = Image(client_logo_path, width=3 * cm, height=1.5 * cm)
            elements.append(logo)
        except Exception:
            pass  # logo inválido — continuar sin él

    elements.append(Paragraph(f"{client_name} — {title}", header_style))
    gen_date = datetime.utcnow().strftime("%d/%m/%Y %H:%M UTC")
    elements.append(Paragraph(f"Período: {period} | Generado: {gen_date}", meta_style))
    elements.append(Spacer(1, 0.5 * cm))

    # --- Tabla de datos ---
    if data:
        headers = list(data[0].keys())
        table_data = [headers]
        for row in data:
            table_data.append([str(row.get(h, "")) for h in headers])

        col_count = len(headers)
        available_width = (page_size[0] - 4 * cm)
        col_width = available_width / col_count

        table = Table(table_data, colWidths=[col_width] * col_count, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(table)

    def add_page_header(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(2 * cm, doc.pagesize[1] - 1.2 * cm, f"{client_name} — {title} | {period}")
        canvas.drawRightString(doc.pagesize[0] - 2 * cm, 1 * cm, f"Página {doc.page}")
        canvas.restoreState()

    doc.build(elements, onFirstPage=add_page_header, onLaterPages=add_page_header)
    return file_path


def _generate_excel(
    data: list[dict[str, Any]],
    title: str,
    period: str,
    client_name: str,
) -> str:
    """Genera Excel con openpyxl. Retorna ruta del archivo temporal."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    tmp_dir = tempfile.mkdtemp(prefix="mind_report_")
    safe_title = "".join(c if c.isalnum() or c in " _-" else "_" for c in title)
    file_path = os.path.join(tmp_dir, f"{safe_title}_{period.replace(' ', '_')}.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]  # Excel limita a 31 caracteres

    # --- Metadata ---
    gen_date = datetime.utcnow().strftime("%d/%m/%Y %H:%M UTC")
    ws["A1"] = client_name
    ws["A1"].font = Font(bold=True, size=14, color="1a1a2e")
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = f"Período: {period}"
    ws["A3"].font = Font(size=10, color="666666")
    ws["A4"] = f"Generado: {gen_date}"
    ws["A4"].font = Font(size=9, color="999999")

    if not data:
        wb.save(file_path)
        return file_path

    # --- Encabezados de columnas ---
    header_row = 6
    headers = list(data[0].keys())
    header_fill = PatternFill("solid", fgColor="1a1a2e")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = border

    # --- Datos ---
    alt_fill = PatternFill("solid", fgColor="F5F5F5")
    for row_idx, row_data in enumerate(data, start=header_row + 1):
        fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row_data.get(header))
            if fill:
                cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    # --- Autoajuste de columnas ---
    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            len(str(header)),
            max((len(str(row.get(header, ""))) for row in data), default=0),
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    wb.save(file_path)
    return file_path
