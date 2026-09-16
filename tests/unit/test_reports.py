"""
Property tests y unitarios para mind/reports/generator.py.
Propiedad 15: La generación de informes produce archivos válidos y completos.
Valida: Requisitos 6.1, 6.2, 6.3
"""
from __future__ import annotations

import os
import pytest
from hypothesis import given, settings, HealthCheck
from hypothesis import strategies as st

from mind.reports.generator import ReportGenerationError, generate_report

# Estrategia para filas de datos de muestra
_row_strategy = st.fixed_dictionaries({
    "nombre": st.text(min_size=1, max_size=20),
    "valor": st.floats(min_value=0, max_value=1_000_000, allow_nan=False, allow_infinity=False),
    "periodo": st.just("2024-01"),
})


class TestGenerateReportProperty15:
    """Propiedad 15: archivos válidos y completos."""

    @given(
        data=st.lists(_row_strategy, min_size=1, max_size=50),
        fmt=st.sampled_from(["pdf", "excel"]),
    )
    @settings(max_examples=30, suppress_health_check=[HealthCheck.too_slow])
    def test_report_file_exists_and_is_readable(self, data, fmt, tmp_path):
        import asyncio
        path = asyncio.get_event_loop().run_until_complete(
            generate_report(data, fmt, title="Test Report", period="2024-Q1")
        )
        assert os.path.exists(path)
        if fmt == "excel":
            import openpyxl
            wb = openpyxl.load_workbook(path)
            assert wb.active is not None
        elif fmt == "pdf":
            with open(path, "rb") as f:
                assert f.read(4) == b"%PDF"
        # Limpieza
        os.unlink(path)

    @pytest.mark.asyncio
    async def test_pdf_contains_client_name(self, tmp_path):
        data = [{"col": "val", "num": 1}]
        path = await generate_report(
            data, "pdf", title="Informe Test", period="2024",
            client_name="AcmeCorp"
        )
        with open(path, "rb") as f:
            content = f.read()
        assert b"AcmeCorp" in content
        os.unlink(path)

    @pytest.mark.asyncio
    async def test_excel_metadata_row_present(self):
        data = [{"producto": "A", "ventas": 100}]
        path = await generate_report(
            data, "excel", title="Ventas", period="Enero 2024",
            client_name="TestClient"
        )
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        assert ws["A1"].value == "TestClient"
        assert ws["A2"].value == "Ventas"
        os.unlink(path)


class TestGenerateReportErrors:

    @pytest.mark.asyncio
    async def test_empty_data_raises_error(self):
        with pytest.raises(ReportGenerationError):
            await generate_report([], "pdf", "Test", "2024")

    @pytest.mark.asyncio
    async def test_invalid_format_raises_error(self):
        with pytest.raises(ReportGenerationError):
            await generate_report([{"a": 1}], "csv", "Test", "2024")  # type: ignore
